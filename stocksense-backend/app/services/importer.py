"""Bulk import products (and opening stock) from CSV or Excel.

- New SKU      -> product is created (with opening stock if a qty is given)
- Existing SKU -> product details are updated (stock is NOT changed; use an adjustment)
- dry_run=True -> nothing is saved, you just get the report (preview)
"""
from __future__ import annotations

import csv
import io
from decimal import Decimal
from decimal import InvalidOperation as DecimalError

from app.engine import StockEngine, StockError
from app.engine.models import to_qty

COLUMNS = ["name", "sku", "category", "uom", "min_qty", "reorder_qty",
           "unit_cost", "initial_stock", "warehouse", "location"]

ALIASES = {
    "product": "name", "product_name": "name", "item": "name", "item_name": "name",
    "code": "sku", "sku_code": "sku", "product_code": "sku", "barcode": "sku",
    "unit": "uom", "unit_of_measure": "uom",
    "min": "min_qty", "minimum": "min_qty", "min_stock": "min_qty", "reorder_point": "min_qty",
    "reorder": "reorder_qty", "reorder_quantity": "reorder_qty",
    "cost": "unit_cost", "price": "unit_cost", "unit_price": "unit_cost",
    "qty": "initial_stock", "quantity": "initial_stock", "stock": "initial_stock",
    "opening_stock": "initial_stock", "on_hand": "initial_stock",
    "rack": "location", "bin": "location",
}
NUMERIC = ("min_qty", "reorder_qty", "unit_cost", "initial_stock")
UPDATABLE = ("name", "category", "uom", "min_qty", "reorder_qty", "unit_cost")
MAX_ROWS = 5000


def _norm_header(h) -> str:
    key = str(h or "").strip().lower().replace(" ", "_").replace("-", "_")
    return ALIASES.get(key, key)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)                       # Excel gives 1001.0 for 1001
    return str(v).strip()


def read_rows(filename: str, content: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
            wb.close()
        except Exception as e:
            raise ValueError(f"Could not read the Excel file: {e}")
    elif name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("Unsupported file type. Upload a .csv or .xlsx file")

    if not rows:
        raise ValueError("The file is empty")
    headers = [_norm_header(h) for h in rows[0]]
    if "name" not in headers or "sku" not in headers:
        raise ValueError("The file needs at least 'name' and 'sku' columns")
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError(f"Too many rows (max {MAX_ROWS})")

    out = []
    for n, values in enumerate(rows[1:], start=2):          # n = row number as seen in Excel
        cells = [_cell(v) for v in values]
        if not any(cells):
            continue                                         # skip blank lines
        row = {h: v for h, v in zip(headers, cells) if h}
        row["_row"] = n
        out.append(row)
    return out


def _parse(row: dict) -> dict:
    name, sku = row.get("name", ""), row.get("sku", "")
    if not name:
        raise ValueError("name is empty")
    if not sku:
        raise ValueError("sku is empty")
    data = {"name": name, "sku": sku}
    for f in ("category", "uom"):
        if row.get(f):
            data[f] = row[f]
    for f in NUMERIC:
        raw = row.get(f, "")
        if raw == "":
            continue
        try:
            val = to_qty(raw.replace(",", ""))
            if not val.is_finite() or val < 0:
                raise ValueError
        except (DecimalError, ValueError):
            raise ValueError(f"{f} '{raw}' must be a number >= 0")
        data[f] = val
    data["warehouse"] = row.get("warehouse", "")
    data["location"] = row.get("location", "")
    return data


def import_products(eng: StockEngine, filename: str, content: bytes, user: str = "system",
                    dry_run: bool = True, create_locations: bool = True) -> dict:
    rows = read_rows(filename, content)                     # ValueError -> 400 in the API
    report = {"dry_run": dry_run, "total_rows": len(rows), "created": 0, "updated": 0,
              "skipped": 0, "errors": [], "warnings": []}

    existing = {p.sku.lower(): p for p in eng.store.products.values()}
    locations = {(l.warehouse.lower(), l.name.lower()): l for l in eng.store.locations.values()}
    planned_locations: set = set()
    seen: set = set()

    for row in rows:
        n = row.pop("_row")
        try:
            data = _parse(row)
            key = data["sku"].lower()
            if key in seen:
                raise ValueError(f"duplicate SKU '{data['sku']}' in this file")
            seen.add(key)

            stock = data.pop("initial_stock", Decimal("0"))
            wh, loc_name = data.pop("warehouse"), data.pop("location")
            product = existing.get(key)

            if product:                                      # ---- update existing
                if stock > 0:
                    report["warnings"].append({"row": n, "message":
                        f"stock ignored for existing product {product.sku}; use an adjustment"})
                changes = {k: v for k, v in data.items()
                           if k in UPDATABLE and getattr(product, k) != v}
                if not changes:
                    report["skipped"] += 1
                    continue
                if "uom" in changes and eng._has_history(product.id):
                    raise ValueError("unit of measure can't change after stock has moved")
                if not dry_run:
                    eng.update_product(product.id, user=user, **changes)
                report["updated"] += 1
                continue

            location_id = None                               # ---- create new
            if stock > 0:
                if not wh:
                    raise ValueError("warehouse is required when a stock qty is given")
                loc_name = loc_name or "Stock"
                lkey = (wh.lower(), loc_name.lower())
                loc = locations.get(lkey)
                if loc is None:
                    if not create_locations:
                        raise ValueError(f"location '{wh} / {loc_name}' does not exist")
                    if lkey not in planned_locations:
                        planned_locations.add(lkey)
                        report["warnings"].append({"row": n, "message":
                            f"new location '{wh} / {loc_name}' "
                            + ("will be created" if dry_run else "was created")})
                    if not dry_run:
                        loc = eng.add_location(loc_name, wh)
                        locations[lkey] = loc
                location_id = loc.id if loc else None

            if not dry_run:
                eng.add_product(**data, initial_stock=stock, location_id=location_id, user=user)
            report["created"] += 1

        except (ValueError, StockError) as e:
            report["errors"].append({"row": n, "message": str(e)})

    return report


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    w.writerow(["Steel Rod", "STL-001", "Raw Material", "kg", "50", "200", "65", "100",
                "Main Warehouse", "Stock"])
    w.writerow(["Office Chair", "CHR-001", "Furniture", "unit", "10", "40", "2500", "", "", ""])
    return buf.getvalue()